"""Pembuatan file Excel rekap anggaran vs realisasi.

Menyediakan workbook: bulanan, tahunan (12 bulan per unit), dan rentang multi-bulan.
Setiap lembar dilengkapi kop bermerek + logo perisai, grafik batang pagu vs
realisasi, dan blok tanda tangan manajemen siap ditandatangani.
"""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage

TEAL = "0D3C45"
TEAL2 = "14758A"
ORANGE = "F2941F"
RED = "DC2626"
GREY = "777777"
MONEY_FMT = "#,##0"
PCT_FMT = "0.0%"

MONTHS_ID = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
             "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

LOGO_PATH = "/app/frontend/public/logo-icon.png"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor=TEAL)
_TOTAL_FILL = PatternFill("solid", fgColor="F5F5F5")
_WHITE_BOLD = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def period_label(period):
    try:
        y, m = period.split("-")
        return f"{MONTHS_ID[int(m) - 1]} {y}"
    except Exception:
        return period


def _kop(ws, subtitle, actor_name, span_cols):
    """Tulis kop (logo di A + teks brand) pada baris 1-4. Return baris awal tabel (6)."""
    last = get_column_letter(max(span_cols, 2))
    ws.merge_cells(f"B1:{last}1")
    ws["B1"] = "PERMINTAAN KEUANGAN"
    ws["B1"].font = Font(bold=True, size=15, color=TEAL)
    ws.merge_cells(f"B2:{last}2")
    ws["B2"] = "SISTEM PENGAJUAN BARANG & JASA · PT. Sumber Berdaya Bersama"
    ws["B2"].font = Font(bold=True, size=9, color=TEAL2)
    ws.merge_cells(f"B3:{last}3")
    ws["B3"] = subtitle
    ws["B3"].font = Font(bold=True, size=11, color=TEAL2)
    ws.merge_cells(f"B4:{last}4")
    ws["B4"] = f"Dicetak: {datetime.now(timezone.utc).strftime('%d-%m-%Y %H:%M')} UTC · oleh {actor_name}"
    ws["B4"].font = Font(size=9, color=GREY)
    for rr in range(1, 5):
        ws.row_dimensions[rr].height = 16
    ws.column_dimensions["A"].width = 10
    try:
        img = XLImage(LOGO_PATH)
        img.width = 60
        img.height = 60
        ws.add_image(img, "A1")
    except Exception:
        pass
    return 6


def _signature_block(ws, start_row, left_col=2, right_col=5):
    r = start_row
    ws.cell(row=r, column=left_col, value="Disiapkan oleh,").font = Font(size=10, bold=True)
    ws.cell(row=r, column=right_col, value="Disetujui oleh,").font = Font(size=10, bold=True)
    ws.cell(row=r + 1, column=left_col, value="Bagian Keuangan").font = Font(size=9, color=GREY)
    ws.cell(row=r + 1, column=right_col, value="Manajemen").font = Font(size=9, color=GREY)
    name_row = r + 5
    line = Border(top=Side(style="thin", color="999999"))
    for col in (left_col, right_col):
        c = ws.cell(row=name_row, column=col, value="(............................)")
        c.font = Font(size=10)
        c.border = line
    ws.cell(row=name_row + 1, column=left_col, value="Tgl: .......................").font = Font(size=9, color=GREY)
    ws.cell(row=name_row + 1, column=right_col, value="Tgl: .......................").font = Font(size=9, color=GREY)
    return name_row + 2


def _bar_chart(ws, title, cat_col, first_data_row, last_data_row, series_cols, header_row):
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.height = 8
    chart.width = 18
    chart.gapWidth = 60
    cats = Reference(ws, min_col=cat_col, min_row=first_data_row, max_row=last_data_row)
    for sc in series_cols:
        ref = Reference(ws, min_col=sc, min_row=header_row, max_row=last_data_row)
        chart.add_data(ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.numFmt = MONEY_FMT
    colors = [TEAL, ORANGE, "94A3B8"]
    for i, s in enumerate(chart.series):
        try:
            s.graphicalProperties.solidFill = colors[i % len(colors)]
        except Exception:
            pass
    return chart


def _header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = _HEADER_FILL
        c.font = _WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER


def _write_monthly_table(ws, data, start_row):
    rows = data.get("rows", [])
    total_pagu = data.get("total_pagu", 0) or 0
    total_real = data.get("total_realisasi", 0) or 0
    total_sisa = total_pagu - total_real
    total_persen = (total_real / total_pagu) if total_pagu else 0

    hr = start_row
    _header_row(ws, hr, ["No", "Unit Kerja", "Pagu Anggaran", "Realisasi", "Sisa", "Serapan %", "Jml Dok"])
    r = hr + 1
    for idx, row in enumerate(rows, start=1):
        pagu = row.get("amount", 0) or 0
        realisasi = row.get("realisasi", 0) or 0
        vals = [idx, row.get("unit_kerja", ""), pagu, realisasi, row.get("sisa", 0) or 0,
                (row.get("persen", 0) or 0) / 100.0, row.get("doc_count", 0) or 0]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i in (3, 4, 5):
                c.number_format = MONEY_FMT
                c.alignment = Alignment(horizontal="right")
            elif i == 6:
                c.number_format = PCT_FMT
                c.alignment = Alignment(horizontal="center")
            elif i in (1, 7):
                c.alignment = Alignment(horizontal="center")
        if pagu > 0 and realisasi > pagu:
            ws.cell(row=r, column=5).font = Font(color=RED, bold=True)
        if row.get("no_budget"):
            ws.cell(row=r, column=2).font = Font(color=ORANGE, italic=True)
        r += 1
    data_end = r - 1
    for i in range(1, 8):
        c = ws.cell(row=r, column=i)
        c.fill = _TOTAL_FILL
        c.border = BORDER
        c.font = Font(bold=True)
    ws.cell(row=r, column=2, value="TOTAL")
    ws.cell(row=r, column=3, value=total_pagu).number_format = MONEY_FMT
    ws.cell(row=r, column=4, value=total_real).number_format = MONEY_FMT
    ws.cell(row=r, column=5, value=total_sisa).number_format = MONEY_FMT
    tp = ws.cell(row=r, column=6, value=total_persen)
    tp.number_format = PCT_FMT
    tp.alignment = Alignment(horizontal="center")
    return {"header_row": hr, "data_start": hr + 1, "data_end": data_end, "total_row": r, "n": len(rows)}


def _apply_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _monthly_sheet(ws, data, period, actor_name, with_kop=True):
    if with_kop:
        start = _kop(ws, f"Rekap Anggaran vs Realisasi · {period_label(period)}", actor_name, 7)
    else:
        ws.merge_cells("A1:G1")
        ws["A1"] = f"Rekap Anggaran vs Realisasi · {period_label(period)}"
        ws["A1"].font = Font(bold=True, size=12, color=TEAL)
        start = 3
    info = _write_monthly_table(ws, data, start)
    _apply_widths(ws, [10, 34, 18, 18, 18, 11, 9])
    ws.freeze_panes = f"A{info['data_start']}"
    if info["n"] > 0:
        chart = _bar_chart(ws, f"Pagu vs Realisasi · {period_label(period)}",
                           cat_col=2, first_data_row=info["data_start"], last_data_row=info["data_end"],
                           series_cols=[3, 4], header_row=info["header_row"])
        ws.add_chart(chart, f"I{start}")
    _signature_block(ws, info["total_row"] + 3)


def build_monthly_workbook(data, period, actor_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Anggaran vs Realisasi"
    _monthly_sheet(ws, data, period, actor_name, with_kop=True)
    return _to_bytes(wb)


def build_range_workbook(months, actor_name):
    """months: list of (period, monthly_data). Lembar Ringkasan + satu lembar per bulan."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    start_label = period_label(months[0][0]) if months else ""
    end_label = period_label(months[-1][0]) if months else ""
    start = _kop(ws, f"Rekap Anggaran vs Realisasi · {start_label} s/d {end_label}", actor_name, 6)
    _header_row(ws, start, ["No", "Periode", "Total Pagu", "Total Realisasi", "Sisa", "Serapan %"])
    r = start + 1
    g_pagu = g_real = 0.0
    for idx, (period, data) in enumerate(months, start=1):
        pagu = data.get("total_pagu", 0) or 0
        real = data.get("total_realisasi", 0) or 0
        g_pagu += pagu
        g_real += real
        vals = [idx, period_label(period), pagu, real, pagu - real, (real / pagu) if pagu else 0]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i in (3, 4, 5):
                c.number_format = MONEY_FMT
                c.alignment = Alignment(horizontal="right")
            elif i == 6:
                c.number_format = PCT_FMT
                c.alignment = Alignment(horizontal="center")
            elif i == 1:
                c.alignment = Alignment(horizontal="center")
        if pagu > 0 and real > pagu:
            ws.cell(row=r, column=5).font = Font(color=RED, bold=True)
        r += 1
    data_end = r - 1
    for i in range(1, 7):
        c = ws.cell(row=r, column=i)
        c.fill = _TOTAL_FILL
        c.border = BORDER
        c.font = Font(bold=True)
    ws.cell(row=r, column=2, value="TOTAL")
    ws.cell(row=r, column=3, value=g_pagu).number_format = MONEY_FMT
    ws.cell(row=r, column=4, value=g_real).number_format = MONEY_FMT
    ws.cell(row=r, column=5, value=g_pagu - g_real).number_format = MONEY_FMT
    tp = ws.cell(row=r, column=6, value=(g_real / g_pagu) if g_pagu else 0)
    tp.number_format = PCT_FMT
    tp.alignment = Alignment(horizontal="center")
    _apply_widths(ws, [6, 26, 20, 20, 20, 12])
    ws.freeze_panes = f"A{start + 1}"
    if data_end >= start + 1:
        chart = _bar_chart(ws, "Pagu vs Realisasi per Periode", cat_col=2,
                           first_data_row=start + 1, last_data_row=data_end,
                           series_cols=[3, 4], header_row=start)
        ws.add_chart(chart, f"H{start}")
    _signature_block(ws, r + 3)

    for period, data in months:
        title = period[:31]
        detail = wb.create_sheet(title=title)
        _monthly_sheet(detail, data, period, actor_name, with_kop=True)
    return _to_bytes(wb)


def build_annual_workbook(annual, year, unit_kerja, actor_name):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Tahunan {year}"[:31]
    sub = f"Rekap Tahunan Anggaran vs Realisasi {year}"
    if unit_kerja:
        sub += f" · Unit: {unit_kerja}"
    start = _kop(ws, sub, actor_name, 6)

    # --- Bagian A: Ringkasan per Unit Kerja ---
    ws.cell(row=start, column=1, value="A. Ringkasan per Unit Kerja").font = Font(bold=True, size=11, color=TEAL)
    hrA = start + 1
    _header_row(ws, hrA, ["No", "Unit Kerja", "Pagu Setahun", "Realisasi Setahun", "Sisa", "Serapan %"])
    r = hrA + 1
    per_unit = annual.get("per_unit", [])
    for idx, u in enumerate(per_unit, start=1):
        pagu = u.get("pagu", 0) or 0
        real = u.get("realisasi", 0) or 0
        vals = [idx, u.get("unit_kerja", ""), pagu, real, pagu - real, (real / pagu) if pagu else 0]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i in (3, 4, 5):
                c.number_format = MONEY_FMT
                c.alignment = Alignment(horizontal="right")
            elif i == 6:
                c.number_format = PCT_FMT
                c.alignment = Alignment(horizontal="center")
            elif i == 1:
                c.alignment = Alignment(horizontal="center")
        if pagu > 0 and real > pagu:
            ws.cell(row=r, column=5).font = Font(color=RED, bold=True)
        r += 1
    for i in range(1, 7):
        c = ws.cell(row=r, column=i)
        c.fill = _TOTAL_FILL
        c.border = BORDER
        c.font = Font(bold=True)
    tot_pagu = annual.get("total_pagu", 0) or 0
    tot_real = annual.get("total_realisasi", 0) or 0
    ws.cell(row=r, column=2, value="TOTAL")
    ws.cell(row=r, column=3, value=tot_pagu).number_format = MONEY_FMT
    ws.cell(row=r, column=4, value=tot_real).number_format = MONEY_FMT
    ws.cell(row=r, column=5, value=tot_pagu - tot_real).number_format = MONEY_FMT
    tpA = ws.cell(row=r, column=6, value=(tot_real / tot_pagu) if tot_pagu else 0)
    tpA.number_format = PCT_FMT
    tpA.alignment = Alignment(horizontal="center")

    # --- Bagian B: Pagu vs Realisasi per Bulan ---
    b_title_row = r + 2
    ws.cell(row=b_title_row, column=1, value="B. Pagu vs Realisasi per Bulan").font = Font(bold=True, size=11, color=TEAL)
    hrB = b_title_row + 1
    _header_row(ws, hrB, ["Bulan", "Pagu", "Realisasi", "Sisa", "Serapan %"])
    rb = hrB + 1
    for m in annual.get("per_month", []):
        pagu = m.get("pagu", 0) or 0
        real = m.get("realisasi", 0) or 0
        vals = [MONTHS_ID[m["month"] - 1], pagu, real, pagu - real, (real / pagu) if pagu else 0]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=rb, column=i, value=v)
            c.border = BORDER
            if i in (2, 3, 4):
                c.number_format = MONEY_FMT
                c.alignment = Alignment(horizontal="right")
            elif i == 5:
                c.number_format = PCT_FMT
                c.alignment = Alignment(horizontal="center")
        if pagu > 0 and real > pagu:
            ws.cell(row=rb, column=4).font = Font(color=RED, bold=True)
        rb += 1
    b_data_end = rb - 1
    for i in range(1, 6):
        c = ws.cell(row=rb, column=i)
        c.fill = _TOTAL_FILL
        c.border = BORDER
        c.font = Font(bold=True)
    ws.cell(row=rb, column=1, value="TOTAL")
    ws.cell(row=rb, column=2, value=tot_pagu).number_format = MONEY_FMT
    ws.cell(row=rb, column=3, value=tot_real).number_format = MONEY_FMT
    ws.cell(row=rb, column=4, value=tot_pagu - tot_real).number_format = MONEY_FMT
    tpB = ws.cell(row=rb, column=5, value=(tot_real / tot_pagu) if tot_pagu else 0)
    tpB.number_format = PCT_FMT
    tpB.alignment = Alignment(horizontal="center")

    _apply_widths(ws, [10, 30, 18, 18, 18, 11])
    ws.freeze_panes = f"A{hrA + 1}"

    chart = _bar_chart(ws, f"Tren Pagu vs Realisasi {year}", cat_col=1,
                       first_data_row=hrB + 1, last_data_row=b_data_end,
                       series_cols=[2, 3], header_row=hrB)
    ws.add_chart(chart, f"H{hrA}")

    _signature_block(ws, rb + 3)
    return _to_bytes(wb)


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def month_range(start, end, limit=12):
    """Kembalikan daftar periode 'YYYY-MM' dari start s/d end (inklusif), maks `limit`."""
    try:
        ys, ms = (int(x) for x in start.split("-"))
        ye, me = (int(x) for x in end.split("-"))
    except Exception:
        return []
    if not (1 <= ms <= 12 and 1 <= me <= 12):
        return []
    if (ys, ms) > (ye, me):
        ys, ms, ye, me = ye, me, ys, ms
    out = []
    cy, cm = ys, ms
    while (cy, cm) <= (ye, me) and len(out) < limit:
        out.append(f"{cy:04d}-{cm:02d}")
        cm += 1
        if cm > 12:
            cm = 1
            cy += 1
    return out
